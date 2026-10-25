"""Build student/teacher DOCX and PDF from one validated UTF-8 JSON master."""
import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata

VERSION = "0.1.2"
LABELS = ("A", "B", "C", "D")

def require(test, message):
    if not test:
        raise ValueError(message)

def text(value, label):
    require(isinstance(value, str) and bool(value.strip()), label + ": expected nonempty text")
    require("\ufffd" not in value, label + ": replacement character found")
    return value

def number(value, label):
    require(isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value) and value > 0, label + ": expected a positive number")
    return value

def strings(values, label):
    require(isinstance(values, list) and bool(values), label + ": expected a nonempty list")
    for value in values:
        text(value, label)

def blocks(value, base):
    if isinstance(value, str):
        text(value, "paragraph")
        return
    require(isinstance(value, list), "content blocks must be a list")
    for block in value:
        require(isinstance(block, dict), "block must be an object")
        kind = block.get("type")
        if kind == "paragraph":
            text(block.get("text"), "paragraph.text")
        elif kind == "table":
            strings(block.get("headers"), "table.headers")
            require(isinstance(block.get("rows"), list) and bool(block["rows"]), "table.rows required")
            for row in block["rows"]:
                strings(row, "table.row")
                require(len(row) == len(block["headers"]), "table row width mismatch")
            text(block.get("caption"), "table.caption")
        elif kind == "image":
            text(block.get("path"), "image.path")
            require(not re.match(r"^[a-zA-Z]+://", block["path"]), "remote image not supported")
            require((base / block["path"]).is_file(), "missing image: " + block["path"])
            width = block.get("width_cm", 12)
            number(width, "image.width_cm")
            require(width <= 16.5, "image wider than printable area")
            text(block.get("caption"), "image.caption")
        else:
            raise ValueError("unsupported block: " + str(kind))

def validate(data, base):
    for key in ("title", "subject", "grade", "unit", "version", "difficulty"):
        text(data.get(key), key)
    number(data.get("minutes"), "minutes")
    number(data.get("total_score"), "total_score")
    strings(data.get("instructions"), "instructions")
    blocks(data.get("materials", []), base)
    questions = data.get("questions")
    require(isinstance(questions, list) and bool(questions), "questions required")
    require([q.get("id") for q in questions] == list(range(1, len(questions) + 1)), "question ids must be sequential")
    seen, previous = set(), None
    for q in questions:
        number(q.get("score"), "question.score")
        blocks(q.get("materials", []), base)
        require(bool(q.get("stem")), "stem required")
        blocks(q["stem"], base)
        text(q.get("answer"), "answer")
        text(q.get("locator"), "locator")
        text(q.get("misconception"), "misconception")
        strings(q.get("explanation"), "explanation")
        group = q.get("group")
        if group:
            text(group, "group")
            require(group == previous or group not in seen, "group questions must be contiguous")
            seen.add(group)
        previous = group
        if q.get("type") == "single":
            for key in ("options", "option_explanations"):
                require(isinstance(q.get(key), dict) and set(q[key]) == set(LABELS), key + ": A-D required")
                for v in q[key].values():
                    text(v, key)
            require(q["answer"] in LABELS, "single answer must be A-D")
        elif q.get("type") == "short":
            lines = q.get("answer_lines", 4)
            require(isinstance(lines, int) and not isinstance(lines, bool) and lines > 0, "answer_lines must be a positive integer")
            strings(q.get("accepted_answers"), "accepted_answers")
            rubric = q.get("rubric")
            require(isinstance(rubric, list) and bool(rubric), "rubric required")
            for row in rubric:
                text(row.get("criterion"), "rubric.criterion")
                number(row.get("points"), "rubric.points")
            require(math.isclose(sum(r["points"] for r in rubric), q["score"]), "rubric points mismatch")
            examples = q.get("partial_credit_examples")
            require(isinstance(examples, list) and bool(examples), "partial_credit_examples required")
            for row in examples:
                text(row.get("response"), "partial.response")
                text(row.get("reason"), "partial.reason")
                number(row.get("points"), "partial.points")
                require(row["points"] < q["score"], "partial credit example must be less than full score")
        else:
            raise ValueError("unsupported question type")
    require(math.isclose(sum(q["score"] for q in questions), data["total_score"]), "total score mismatch")
    blueprint = data.get("blueprint")
    require(isinstance(blueprint, list) and [b.get("id") for b in blueprint] == [q["id"] for q in questions],
            "blueprint must contain every question once in order")
    for row in blueprint:
        for key in ("locator", "goal", "task", "basis"):
            text(row.get(key), "blueprint." + key)
        number(row.get("minutes"), "blueprint.minutes")
    require(math.isclose(sum(b["minutes"] for b in blueprint), data["minutes"]), "estimated minutes mismatch")
    sources = data.get("sources")
    require(isinstance(sources, list) and bool(sources), "sources required (may record unverified status)")
    statuses = {"已讀正文", "僅摘要／目錄", "無法讀取", "尚未核對", "技能內既有查讀摘要，非本次連線讀取"}
    for source in sources:
        for key in ("institution", "document", "year", "accessed", "locator", "used_for", "status", "kind"):
            text(source.get(key), "sources." + key)
        date.fromisoformat(source["accessed"])
        require(source["kind"] in ("official", "material"), "source.kind must be official/material")
        require(source["status"] in statuses, "unsupported source.status")
        url = source.get("url", "")
        require(isinstance(url, str), "source.url must be a string")
        if url or source["kind"] == "official":
            require(bool(re.match(r"^https?://[^\s/]+(?:/[^\s]*)?$", url)), "invalid source URL")
    return data

def normal(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value))

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def safe_name(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:80] or "assessment"

class Paper:
    def __init__(self, data, role, base, font, size):
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        self.Cm, self.Pt, self.qn, self.Element = Cm, Pt, qn, OxmlElement
        self.doc, self.expected, self.images, self.base = Document(), [], [], base
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
        for name, points in (("Normal", size), ("Title", 20), ("Heading 1", 16), ("Heading 2", 14)):
            style = self.doc.styles[name]
            style.font.name, style.font.size, style.font.color.rgb = font, Pt(points), RGBColor(0, 0, 0)
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ("eastAsia", "ascii", "hAnsi"):
                fonts.set(qn("w:" + attr), font)
            style.paragraph_format.space_after = Pt(6)
        self.doc.styles["Normal"].paragraph_format.line_spacing = 1.15
        self.doc.styles["Normal"].paragraph_format.space_after = Pt(4)
        # Remove inherited title rules and other template paragraph borders.
        for style in self.doc.styles:
            ppr = style.element.find(qn("w:pPr"))
            if ppr is not None:
                for border in list(ppr.findall(qn("w:pBdr"))):
                    ppr.remove(border)
        self.doc.core_properties.title = data["title"] + " " + role
        self.doc.core_properties.author = ""
        footer = section.footer.paragraphs[0]
        footer.alignment = 1
        footer.add_run("第 ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        footer.add_run(" 頁")
        for run in footer.runs:
            run.font.name, run.font.size = font, Pt(10)
        self.p(data["title"] + " " + data["version"] + " " + role, "Title")
        self.p(f'{data["subject"]}  {data["grade"]}  {data["unit"]}')
        self.p(f'作答時間 {data["minutes"]:g} 分鐘    總分 {data["total_score"]:g} 分    共 {len(data["questions"])} 小題')
        self.p("班級＿＿＿＿  姓名＿＿＿＿  座號＿＿＿＿")
        for instruction in data["instructions"]:
            self.p(instruction)
        self.content(data.get("materials", []))
        for q in data["questions"]:
            block_start = len(self.doc.element.body) - 1
            self.content(q.get("materials", []))
            title = f'第{q["id"]}題  {q["score"]:g}分'
            if q.get("group"):
                title += "  " + q["group"]
            self.p(title, "Heading 2")
            self.content(q["stem"])
            if q["type"] == "single":
                for index, label in enumerate(LABELS):
                    p = self.p(label + "  " + q["options"][label], keep=index < 3)
                    p.paragraph_format.space_after = Pt(2)
            else:
                for index in range(q.get("answer_lines", 4)):
                    p = self.doc.add_paragraph("＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿")
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.keep_with_next = index < q.get("answer_lines", 4) - 1
                    p.paragraph_format.keep_together = True
            self.keep_block(block_start)

    def keep_block(self, start):
        from docx.text.paragraph import Paragraph
        paragraphs = []
        for child in list(self.doc.element.body)[start:-1]:
            paragraphs.extend([child] if child.tag == self.qn("w:p") else list(child.iter(self.qn("w:p"))))
        for index, node in enumerate(paragraphs):
            Paragraph(node, self.doc._body).paragraph_format.keep_with_next = index < len(paragraphs) - 1

    def p(self, value, style=None, keep=False):
        p = self.doc.add_paragraph(value, style)
        p.paragraph_format.keep_with_next = keep or style in ("Title", "Heading 1", "Heading 2")
        p.paragraph_format.widow_control = True
        p.paragraph_format.keep_together = True
        self.expected.append(value)
        return p

    def table(self, headers, rows):
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        widths = [14.0, 2.5] if headers == ["評分要點", "分數"] else [16.5 / len(headers)] * len(headers)
        for cell, width in zip(table.columns, widths):
            cell.width = self.Cm(width)
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
            self.expected.append(value)
        repeat = self.Element("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = value
                self.expected.append(value)
        borders = self.Element("w:tblBorders")
        for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = self.Element("w:" + name)
            for attr, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                edge.set(self.qn("w:" + attr), val)
            borders.append(edge)
        table._tbl.tblPr.append(borders)
        for index, row in enumerate(table.rows):
            prevent = self.Element("w:cantSplit")
            row._tr.get_or_add_trPr().append(prevent)
            for column_index, cell in enumerate(row.cells):
                cell.width = self.Cm(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if index == 0:
                    shade = self.Element("w:shd")
                    shade.set(self.qn("w:fill"), "EEEEEE")
                    cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = self.Pt(3)
                    p.paragraph_format.space_after = self.Pt(3)
                    for run in p.runs:
                        run.font.size = self.Pt(11)
                        run.bold = index == 0

    def content(self, value):
        if isinstance(value, str):
            self.p(value, keep=True)
            return
        for block in value:
            if block["type"] == "paragraph":
                self.p(block["text"])
            elif block["type"] == "table":
                self.p(block["caption"], keep=True)
                self.table(block["headers"], block["rows"])
            else:
                path = (self.base / block["path"]).resolve()
                self.doc.add_picture(str(path), width=self.Cm(block.get("width_cm", 12)))
                self.images.append(str(path))
                self.p(block["caption"])

    def teacher(self, data):
        self.doc.styles["Normal"].paragraph_format.space_after = self.Pt(2)
        self.doc.styles["Heading 1"].paragraph_format.space_before = self.Pt(0)
        self.p("答案與完整詳解", "Heading 1").paragraph_format.page_break_before = True
        for q in data["questions"]:
            block_start = len(self.doc.element.body) - 1
            heading = self.p(f'第{q["id"]}題答案與評分  {q["score"]:g}分', "Heading 2")
            if q["type"] == "short":
                heading.paragraph_format.page_break_before = True
            self.p("答案：" + q["answer"])
            for paragraph in q["explanation"]:
                self.p(paragraph)
            if q["type"] == "single":
                for label in LABELS:
                    self.p(label + " 選項解析：" + q["option_explanations"][label])
            else:
                self.p("可接受的表達", "Heading 2")
                for answer in q["accepted_answers"]:
                    self.p(answer)
                self.table(["評分要點", "分數"], [[r["criterion"], f'{r["points"]:g}'] for r in q["rubric"]])
                self.p("部分給分例", "Heading 2")
                for example in q["partial_credit_examples"]:
                    self.p(f'回答：{example["response"]}；給 {example["points"]:g} 分。{example["reason"]}')
            self.p("教材定位：" + q["locator"])
            self.p("常見迷思：" + q["misconception"])
            self.keep_block(block_start)
        self.p("命題藍圖", "Heading 1").paragraph_format.page_break_before = True
        self.p("以下作答時間為未施測估計。")
        for b, q in zip(data["blueprint"], data["questions"]):
            block_start = len(self.doc.element.body) - 1
            kind = "單選" if q["type"] == "single" else "簡答"
            self.p(f'第{b["id"]}題  {kind}  {q["score"]:g}分  預估{b["minutes"]:g}分鐘', "Heading 2")
            for key, label in (("locator", "教材定位"), ("goal", "能力與知識"), ("task", "情境與判斷任務"), ("basis", "答案或評分依據")):
                self.p(label + "：" + b[key])
            self.keep_block(block_start)
        self.p("教材與官方依據查讀紀錄", "Heading 1")
        for i, source in enumerate(data["sources"], 1):
            block_start = len(self.doc.element.body) - 1
            self.p(f'來源{i}  {source["document"]}', "Heading 2")
            for key, label in (("institution", "機構或作者"), ("year", "年份"), ("url", "網址"),
                               ("accessed", "查閱日期"), ("locator", "正文定位"),
                               ("used_for", "採用處"), ("status", "讀取狀態"), ("kind", "來源類型")):
                self.p(label + "：" + (source.get(key) or "本機附件，無外部網址"))
            self.keep_block(block_start)

def check_text(expected, actual):
    content = normal(actual)
    # Fragment checks are deterministic completeness aids, not semantic or visual QA.
    missing = [value for value, count in Counter(expected).items() if content.count(normal(value)) < count]
    return {"expected_fragments": len(expected), "missing_fragments": missing, "passed": not missing}

def export_pdf(docx, engine, soffice):
    pdf = docx.with_suffix(".pdf")
    if engine == "auto":
        engine = "libreoffice" if soffice else ("word" if os.name == "nt" else "unavailable")
    if engine == "word":
        require(os.name == "nt", "Word conversion requires Windows")
        helper = Path(__file__).with_name("export_pdf.ps1")
        command = ["powershell.exe", "-NoProfile", "-NonInteractive", "-File", str(helper),
                   "-InputDocx", str(docx), "-OutputPdf", str(pdf)]
        subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120, check=True)
    elif engine == "libreoffice":
        require(soffice and Path(soffice).is_file(), "supply a verified --soffice executable")
        with tempfile.TemporaryDirectory(prefix="assessment-lo-") as temp:
            profile = Path(temp).as_uri()
            command = [str(soffice), "-env:UserInstallation=" + profile, "--headless",
                       "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)]
            subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120, check=True)
    else:
        raise ValueError("No PDF converter available")
    require(pdf.is_file() and pdf.stat().st_size > 0, "PDF export missing or empty")
    return pdf

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--pdf-engine", choices=("auto", "word", "libreoffice", "none"), default="auto")
    parser.add_argument("--soffice", type=Path)
    parser.add_argument("--pdftoppm")
    parser.add_argument("--preview-engine", choices=("auto", "poppler", "pdfium"), default="auto")
    parser.add_argument("--font", default="Microsoft JhengHei")
    parser.add_argument("--font-size", type=float, default=12)
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    source = args.input.resolve()
    data = validate(json.loads(source.read_text(encoding="utf-8-sig")), source.parent)
    if args.validate_only:
        print("Master validated")
        return 0
    require(args.output is not None, "--output is required")
    number(args.font_size, "font-size")
    out = args.output.resolve()
    require(not out.exists(), "output folder already exists; choose a new version folder")
    from docx import Document
    from pypdf import PdfReader
    out.mkdir(parents=True)
    qa = {"script_version": VERSION, "visual_review": "pending", "semantic_review": "pending",
          "documents": {}, "errors": [], "four_files_complete": False}
    # Freeze referenced images beside the retained master so it can be rebuilt elsewhere.
    frozen = json.loads(json.dumps(data, ensure_ascii=False))
    for values in [frozen.get("materials", [])] + [q.get("materials", []) for q in frozen["questions"]] + [q["stem"] for q in frozen["questions"]]:
        if not isinstance(values, list):
            continue
        for block in values:
            if block.get("type") == "image":
                original = (source.parent / block["path"]).resolve()
                asset = out / "assets" / (digest(original)[:12] + "_" + safe_name(original.name))
                asset.parent.mkdir(exist_ok=True)
                shutil.copy2(original, asset)
                block["path"] = asset.relative_to(out).as_posix()
    (out / "master.json").write_text(json.dumps(frozen, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {"source_master": str(source), "source_master_sha256": digest(source),
                "created_at_utc": datetime.now(timezone.utc).isoformat(), "script_version": VERSION,
                "python": sys.version, "packages": {p: importlib.metadata.version(p) for p in ("python-docx", "pypdf", "Pillow", "pypdfium2")},
                "pdf_engine_requested": args.pdf_engine, "font": args.font, "font_size": args.font_size,
                "asset_sha256": {}, "render_dpi": 110}
    stem = safe_name(data["unit"] + "_" + data["version"])
    papers = {}
    for role in ("學生卷", "教師卷"):
        paper = Paper(data, role, source.parent, args.font, args.font_size)
        if role == "教師卷":
            paper.teacher(data)
        docx = out / (stem + "_" + role + ".docx")
        paper.doc.save(docx)
        document = Document(docx)
        body = "".join(node.text or "" for node in document.element.body.iter() if node.tag.endswith("}t"))
        qa["documents"][docx.name] = {"sha256": digest(docx), "text_check": check_text(paper.expected, body),
                                     "image_count": len(paper.images)}
        manifest["asset_sha256"].update({path: digest(Path(path)) for path in paper.images})
        papers[docx] = paper
    if args.pdf_engine != "none":
        for docx, paper in papers.items():
            try:
                pdf = export_pdf(docx, args.pdf_engine, args.soffice)
                reader = PdfReader(pdf)
                # Page-number footers can occur inside a paragraph's extracted text across pages.
                extracted = "\n".join(re.sub(r"(?m)^第\s*\d+\s*頁[ \t]*$", "", page.extract_text() or "") for page in reader.pages)
                record = {"sha256": digest(pdf), "page_count": len(reader.pages),
                          "text_check": check_text(paper.expected, extracted), "rendered_pages": []}
                require(len(reader.pages) > 0, "PDF has no pages")
                qa["documents"][pdf.name] = record
                renderer = None if args.preview_engine == "pdfium" else (args.pdftoppm or shutil.which("pdftoppm"))
                if renderer:
                    preview = out / "preview" / docx.stem
                    preview.mkdir(parents=True)
                    subprocess.run([renderer, "-r", "110", "-png", str(pdf), str(preview / "page")],
                                   capture_output=True, timeout=120, check=True)
                    record["rendered_pages"] = [str(p.relative_to(out)) for p in sorted(preview.glob("page-*.png"))]
                    require(len(record["rendered_pages"]) == len(reader.pages), "rendered page count mismatch")
                else:
                    if args.preview_engine == "poppler":
                        raise ValueError("pdftoppm unavailable")
                    import pypdfium2 as pdfium
                    preview = out / "preview" / docx.stem
                    preview.mkdir(parents=True)
                    pdf_document = pdfium.PdfDocument(str(pdf))
                    try:
                        for page_index in range(len(pdf_document)):
                            page = pdf_document[page_index]
                            bitmap = page.render(scale=110 / 72)
                            rendered = bitmap.to_pil()
                            rendered.save(preview / f"page-{page_index + 1:02}.png")
                            rendered.close()
                            bitmap.close()
                            page.close()
                    finally:
                        pdf_document.close()
                    record["preview_engine"] = "pypdfium2"
                    record["rendered_pages"] = [str(p.relative_to(out)) for p in sorted(preview.glob("page-*.png"))]
                    require(len(record["rendered_pages"]) == len(reader.pages), "rendered page count mismatch")
            except (ValueError, OSError, subprocess.SubprocessError) as e:
                detail = str(e)
                if isinstance(e, subprocess.CalledProcessError):
                    detail += " " + str(e.stderr)
                qa["errors"].append({"file": docx.name, "error": detail})
    else:
        qa["errors"].append({"stage": "pdf", "error": "PDF intentionally skipped; only DOCX completed"})
    qa["four_files_complete"] = len(qa["documents"]) == 4 and not qa["errors"]
    qa["all_text_checks_passed"] = all(d["text_check"]["passed"] for d in qa["documents"].values())
    (out / "QA.json").write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "build_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(out), "four_files_complete": qa["four_files_complete"],
                      "all_text_checks_passed": qa["all_text_checks_passed"],
                      "visual_review": "pending", "errors": qa["errors"]}, ensure_ascii=False))
    if args.pdf_engine == "none":
        return 0 if qa["all_text_checks_passed"] else 2
    return 0 if qa["four_files_complete"] and qa["all_text_checks_passed"] else 2

if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (ValueError, OSError, json.JSONDecodeError) as error:
        print("ERROR: " + str(error), file=sys.stderr)
        raise SystemExit(2)
